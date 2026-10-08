"""Document AI pipeline: Parser/OCR -> Cleaning -> Chunking -> Embedding -> pgvector."""
import csv
import io
import logging
import os
import re
from typing import List, Tuple

from sqlalchemy import delete, select
from sqlalchemy.orm import Session

from app.core.config import settings
from app.core.database import SessionLocal
from app.models import Document, DocumentChunk
from app.services.embedding_service import embeddings
from app.utils.chunking import chunk_pages

log = logging.getLogger("documents")
ALLOWED_EXT = {"pdf", "docx", "xlsx", "csv", "txt", "png", "jpg", "jpeg"}
IMAGE_EXT = {"png", "jpg", "jpeg"}
_MAGIC = {"pdf": [b"%PDF"], "png": [b"\x89PNG"], "jpg": [b"\xff\xd8"], "jpeg": [b"\xff\xd8"],
          "docx": [b"PK"], "xlsx": [b"PK"]}


def valid_signature(ext: str, head: bytes) -> bool:
    sigs = _MAGIC.get(ext)
    return True if not sigs else any(head.startswith(s) for s in sigs)


def clean_text(t: str) -> str:
    t = t.replace("\x00", " ").replace("\r", "")
    t = re.sub(r"-\n(?=[a-z])", "", t)      # sambung kata yang terpotong di akhir baris
    t = re.sub(r"[ \t]{2,}", " ", t)
    return re.sub(r"\n{3,}", "\n\n", t).strip()


def _group(lines: List[str], per_page: int) -> List[Tuple[int, str]]:
    return [(i // per_page + 1, "\n".join(lines[i:i + per_page])) for i in range(0, len(lines), per_page)]


def extract_pages(path: str, ext: str) -> Tuple[List[Tuple[int, str]], bool]:
    """Return ([(halaman, teks)], used_ocr)"""
    used_ocr = False
    if ext == "pdf":
        import fitz
        from PIL import Image
        from app.services.ocr_service import ocr_image
        pages = []
        with fitz.open(path) as pdf:
            for i, page in enumerate(pdf, 1):
                txt = page.get_text("text").strip()
                if len(txt) < 30:  # kemungkinan PDF hasil scan -> OCR
                    pix = page.get_pixmap(dpi=200)
                    txt = ocr_image(Image.frombytes("RGB", (pix.width, pix.height), pix.samples))
                    used_ocr = True
                pages.append((i, txt))
        return pages, used_ocr
    if ext in IMAGE_EXT:
        from PIL import Image
        from app.services.ocr_service import ocr_image
        return [(1, ocr_image(Image.open(path)))], True
    if ext == "docx":
        import docx
        d = docx.Document(path)
        lines = [p.text for p in d.paragraphs if p.text.strip()]
        for tb in d.tables:
            for row in tb.rows:
                lines.append(" | ".join(c.text.strip() for c in row.cells))
        return _group(lines, 40), False
    if ext == "xlsx":
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        pages = []
        for n, ws in enumerate(wb.worksheets, 1):
            rows = [f"Sheet {ws.title}"] + [" | ".join("" if c is None else str(c) for c in r)
                                            for r in ws.iter_rows(values_only=True, max_row=5000)]
            pages += [(n * 1000 + p, t) for p, t in _group(rows, 50)]
        return [(i + 1, t) for i, (_, t) in enumerate(pages)], False
    if ext == "csv":
        with open(path, encoding="utf-8", errors="ignore", newline="") as f:
            rows = [" | ".join(r) for r in csv.reader(f)][:5000]
        return _group(rows, 50), False
    # txt / markdown: penanda halaman "--- Halaman N ---" atau form-feed
    with open(path, encoding="utf-8", errors="ignore") as f:
        raw = f.read()
    parts = re.split(r"^\s*-{2,}\s*Halaman\s+\d+\s*-{2,}\s*$|\f", raw, flags=re.M | re.I)
    if len(parts) > 1:
        return [(i + 1, p) for i, p in enumerate(parts) if p.strip()], False
    return [(i + 1, raw[j:j + 3000]) for i, j in enumerate(range(0, len(raw), 3000))], False


def ingest_document(doc_id: int) -> None:
    """Dijalankan di background (BackgroundTasks / Celery worker) dengan session sendiri."""
    with SessionLocal() as db:
        doc = db.get(Document, doc_id)
        if not doc:
            return
        try:
            doc.status, doc.error = "PROCESSING", None
            db.commit()
            db.execute(delete(DocumentChunk).where(DocumentChunk.document_id == doc_id))
            pages, used_ocr = extract_pages(doc.file_path, doc.file_type)
            pages = [(n, clean_text(t)) for n, t in pages if t and t.strip()]
            chunks = chunk_pages(pages, settings.chunk_size, settings.chunk_overlap)
            if not chunks:
                raise ValueError("Tidak ada teks yang dapat diekstrak dari dokumen")
            vectors = embeddings.embed_texts([c["content"] for c in chunks])
            for c, v in zip(chunks, vectors):
                db.add(DocumentChunk(document_id=doc_id, chunk_index=c["chunk_index"], content=c["content"],
                                     page_number=c["page_number"], section=c["section"], embedding=v,
                                     meta={"filename": doc.filename, "file_type": doc.file_type}))
            doc.status, doc.used_ocr = "READY", used_ocr
            doc.page_count, doc.chunk_count = len(pages), len(chunks)
            db.commit()
            log.info("Dokumen %s siap: %d halaman, %d chunk", doc.filename, len(pages), len(chunks))
        except Exception as e:  # noqa
            log.exception("Ingest gagal untuk dokumen %s", doc_id)
            db.rollback()
            doc = db.get(Document, doc_id)
            doc.status, doc.error = "FAILED", str(e)[:500]
            db.commit()


def enqueue_ingest(doc_id: int, background_tasks=None) -> None:
    if settings.use_celery:
        from app.worker import ingest_task
        ingest_task.delay(doc_id)
    elif background_tasks is not None:
        background_tasks.add_task(ingest_document, doc_id)
    else:
        ingest_document(doc_id)


def reindex_all() -> int:
    with SessionLocal() as db:
        ids = [i for (i,) in db.execute(select(Document.id)).all()]
    for i in ids:
        ingest_document(i)
    return len(ids)


def delete_file(path: str) -> None:
    try:
        if path and os.path.isfile(path) and os.path.abspath(path).startswith(os.path.abspath(settings.upload_dir)):
            os.remove(path)
    except OSError:
        pass


# ---------- ekstraksi invoice terstruktur ----------
def parse_idr(s: str) -> int:
    digits = re.sub(r"[^\d]", "", s.split(",")[0] if re.search(r",\d{1,2}$", s) else s)
    return int(digits) if digits else 0


def regex_invoice(text: str) -> dict:
    num = re.search(r"(?:invoice|inv|no\.?\s*faktur|nomor)\s*(?:no\.?|number|#|:)?\s*([A-Z]{2,4}[-/ ]?\d[\w\-/]*)", text, re.I)
    date = re.search(r"(\d{4}-\d{2}-\d{2}|\d{1,2}[/-]\d{1,2}[/-]\d{4})", text)
    total = re.findall(r"(?:grand\s*total|total(?:\s*bayar|\s*tagihan)?|jumlah)\s*[:\-]?\s*(?:rp\.?|idr)?\s*([\d\.,]{4,})", text, re.I)
    vendor = next((l.strip() for l in text.splitlines() if len(l.strip()) > 3 and not re.search(r"invoice|faktur", l, re.I)), "")
    return {"invoice_number": num.group(1).strip() if num else None, "vendor": vendor or None,
            "date": date.group(1) if date else None, "total": parse_idr(total[-1]) if total else None}
